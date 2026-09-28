import openpyxl

wb = openpyxl.load_workbook('./cuenta_corriente_franco.xlsx', data_only=True)
sheets_summary = []
for name in wb.sheetnames:
    ws = wb[name]
    sales_count = 0
    payments_count = 0
    last_saldo = 0
    for r in range(2, min(500, ws.max_row + 1)):
        c_val = ws.cell(r, 2).value
        if c_val:
            c_str = str(c_val).strip().upper()
            if 'VENTA' in c_str:
                sales_count += 1
            elif 'PAGO' in c_str:
                payments_count += 1
        saldo_val = ws.cell(r, 21).value # Col U
        if saldo_val is not None:
            try:
                last_saldo = float(saldo_val)
            except:
                pass
    sheets_summary.append((name, sales_count, payments_count, last_saldo))

print('Total sheets:', len(sheets_summary))
sheets_summary.sort(key=lambda x: x[3], reverse=True)
print('\nTop 20 Clients by Debt in Excel:')
for s in sheets_summary[:20]:
    print(f'Client: {s[0]:<20} Sales: {s[1]:<4} Payments: {s[2]:<4} Balance: ${s[3]:,.2f}')
