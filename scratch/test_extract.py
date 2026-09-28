import openpyxl
import json
import datetime
import re

wb = openpyxl.load_workbook('./cuenta_corriente_franco.xlsx', data_only=True)

all_clients = []
all_ventas = []
all_cobranzas = []

vta_counter = 1
cob_counter = 1

def safe_parse_date(f_val, default="2024-09-15"):
    if not f_val:
        return default
    if isinstance(f_val, (datetime.datetime, datetime.date)):
        return f_val.strftime("%Y-%m-%d")
    s = str(f_val).strip()
    match = re.search(r'(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?', s)
    if match:
        day = int(match.group(1))
        month = int(match.group(2))
        yr = int(match.group(3)) if match.group(3) else 2024
        if yr < 100: yr += 2000
        # sanitize month/day
        if 1 <= month <= 12 and 1 <= day <= 31:
            return f"{yr:04d}-{month:02d}-{day:02d}"
    return default

for sheet_idx, sheet_name in enumerate(wb.sheetnames):
    ws = wb[sheet_name]
    cli_id = f"CLI-F-{sheet_idx+1:03d}"
    client_name = str(sheet_name).strip()
    
    header_map = {}
    for c in range(1, min(26, ws.max_column + 1)):
        val = ws.cell(1, c).value
        if val:
            header_map[str(val).strip().upper()] = c
    
    c_fecha = 1
    c_concepto = 2
    c_corte = 3
    c_cortes = header_map.get('CANTIDAD DE CORTES')
    c_totalkg = header_map.get('TOTAL KG')
    c_precio = header_map.get('PRECIO')
    c_debe = header_map.get('DEBE')
    c_pago = header_map.get('PAGO')
    c_transf = header_map.get('TRANSFERENCIA')
    c_saldo = header_map.get('SALDO')
    
    if not c_totalkg: c_totalkg = 15
    if not c_cortes: c_cortes = c_totalkg - 1
    if not c_precio: c_precio = c_totalkg + 1
    if not c_debe: c_debe = c_totalkg + 3
    if not c_pago: c_pago = c_totalkg + 4
    if not c_transf: c_transf = c_totalkg + 5
    if not c_saldo: c_saldo = c_totalkg + 6
        
    client_ventas = []
    client_cobranzas = []
    
    for r in range(2, min(300, ws.max_row + 1)):
        f_val = ws.cell(r, c_fecha).value
        concepto_raw = ws.cell(r, c_concepto).value
        if not f_val and not concepto_raw:
            continue
        
        fecha_str = safe_parse_date(f_val)
        concepto_val = str(concepto_raw or '').strip().upper()
        
        if 'VENTA' in concepto_val:
            corte = str(ws.cell(r, c_corte).value or '1/2 RES').strip()
            pesos = []
            for col_k in range(4, c_cortes):
                w = ws.cell(r, col_k).value
                if w is not None and isinstance(w, (int, float)) and w > 0:
                    pesos.append(float(w))
            
            totalkg_val = ws.cell(r, c_totalkg).value
            try:
                totalkg = float(totalkg_val) if (totalkg_val is not None and str(totalkg_val).strip() != '') else sum(pesos)
            except:
                totalkg = sum(pesos)
            
            cant_cortes = len(pesos) if pesos else 1
            
            precio_val = ws.cell(r, c_precio).value
            try:
                precio = float(precio_val) if (precio_val is not None and str(precio_val).strip() != '') else 6500.0
            except:
                precio = 6500.0
                
            debe_val = ws.cell(r, c_debe).value
            try:
                total_venta = float(debe_val) if (debe_val is not None and str(debe_val).strip() != '') else (totalkg * precio)
            except:
                total_venta = totalkg * precio
                
            if total_venta > 0:
                vta_id = f"VTA-{vta_counter:04d}"
                vta_num = f"VTA-{vta_counter:06d}"
                vta_counter += 1
                client_ventas.append({
                    "id": vta_id,
                    "numero": vta_num,
                    "clienteId": cli_id,
                    "vendedorId": "VEND-FRANCO",
                    "fechaEmision": fecha_str,
                    "corte": corte,
                    "pesos": pesos,
                    "cantidadCortes": cant_cortes,
                    "totalKg": round(totalkg, 2),
                    "precioKg": round(precio, 2),
                    "total": round(total_venta, 2),
                    "saldoPendiente": round(total_venta, 2),
                    "montoCobrado": 0,
                    "estado": "pendiente",
                    "pagosAplicados": []
                })
        elif 'PAGO' in concepto_val:
            pago_efectivo = 0.0
            pago_transf = 0.0
            
            pe = ws.cell(r, c_pago).value
            if pe and isinstance(pe, (int, float)):
                pago_efectivo = float(pe)
            pt = ws.cell(r, c_transf).value
            if pt and isinstance(pt, (int, float)):
                pago_transf = float(pt)
                
            total_cobro = pago_efectivo + pago_transf
            if total_cobro > 0:
                cob_id = f"COB-{cob_counter:04d}"
                cob_num = f"REC-{cob_counter:06d}"
                cob_counter += 1
                metodo = "Mixto (Efectivo y Transf.)" if (pago_efectivo > 0 and pago_transf > 0) else ("Transferencia" if pago_transf > 0 else "Efectivo")
                client_cobranzas.append({
                    "id": cob_id,
                    "numero": cob_num,
                    "clienteId": cli_id,
                    "vendedorId": "VEND-FRANCO",
                    "fecha": fecha_str,
                    "totalCobrado": round(total_cobro, 2),
                    "efectivo": round(pago_efectivo, 2),
                    "transferencia": round(pago_transf, 2),
                    "metodo": metodo,
                    "observaciones": f"Pago Cta Cte ({metodo})"
                })
    
    total_v = sum(v["total"] for v in client_ventas)
    total_c = sum(c["totalCobrado"] for c in client_cobranzas)
    saldo_actual = total_v - total_c
    
    all_clients.append({
        "id": cli_id,
        "razonSocial": client_name,
        "vendedorId": "VEND-FRANCO",
        "telefono": "+54 9 11 4455-8899",
        "contactoNombre": f"Responsable {client_name}",
        "direccion": f"Av. Carnicería Central {sheet_idx+100}, Buenos Aires",
        "limiteCredito": max(5000000.0, round(saldo_actual * 1.5, -5) if saldo_actual > 0 else 5000000.0),
        "plazoDias": 14,
        "estado": "activo",
        "fechaAlta": "2024-01-15"
    })
    all_ventas.extend(client_ventas)
    all_cobranzas.extend(client_cobranzas)

print(f"Extracted successfully: {len(all_clients)} clients, {len(all_ventas)} sales, {len(all_cobranzas)} payments")
