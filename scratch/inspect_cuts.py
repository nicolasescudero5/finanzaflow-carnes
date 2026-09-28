import openpyxl
import datetime

wb = openpyxl.load_workbook('./cuenta_corriente_franco.xlsx', data_only=True)

print("Total client sheets:", len(wb.sheetnames))
cut_types = set()
for name in wb.sheetnames[:20]:
    ws = wb[name]
    for r in range(2, min(50, ws.max_row + 1)):
        corte = ws.cell(r, 3).value
        if corte:
            cut_types.add(str(corte).strip())

print("Sample cuts found in Excel:", sorted(list(cut_types)))
